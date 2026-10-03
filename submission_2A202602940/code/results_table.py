"""results_table.py — Nhiệm vụ: lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx (đừng gõ tay hàng chục dòng, rất dễ sai).

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính, đừng ghi đè)
"""
from __future__ import annotations

import json
import math
from pathlib import Path


def _clean(v):
    """Chuyển sang kiểu JSON chuẩn: tuple -> list, NaN/inf -> None, số numpy/torch -> số Python."""
    if isinstance(v, dict):
        return {k: _clean(x) for k, x in v.items()}
    if isinstance(v, (list, tuple)):
        return [_clean(x) for x in v]
    if hasattr(v, "item") and not isinstance(v, (str, bytes)):
        v = v.item()
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] (KHÔNG ghi best_state) ra
    <results_dir>/<exp_id>.json. Trả về đường dẫn file. Tạo thư mục nếu chưa có.
    NaN/inf (lần chạy phân kỳ) được ghi thành null để file là JSON hợp lệ."""
    out = Path(results_dir)
    out.mkdir(parents=True, exist_ok=True)
    payload = {k: _clean(result[k]) for k in ("cfg", "history", "summary")}
    path = out / f"{result['cfg']['exp_id']}.json"
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=1)
    return str(path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    results = []
    for p in sorted(Path(results_dir).glob("*.json")):
        with open(p, encoding="utf-8") as f:
            results.append(json.load(f))
    return sorted(results, key=lambda r: r["cfg"]["exp_id"])


# Giá trị hợp lệ của các cột có danh sách chọn trong mẫu (Data Validation)
GROUPS = ("baseline", "loss", "optimizer", "hparam", "dropout", "clipping", "amp", "init", "final", "other")
GROUP_ALIAS = {"baseline-lr": "hparam"}     # các lần dò lr cho baseline: thuộc chủ đề hyper-parameter (lr)
LOSS_NAMES = {"ce": "CE", "mse": "MSE"}
OPT_NAMES = {"sgd": "SGD", "sgd_momentum": "SGD+momentum", "adam": "Adam", "adamw": "AdamW"}
FORMULA_COLS = ("step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise")
N_TEMPLATE_ROWS = 60                          # mẫu có công thức sẵn cho dòng 2..61


def _num(v):
    return None if v is None or (isinstance(v, float) and not math.isfinite(v)) else v


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary (+ eval_acc, eval_macro_f1 nếu có)
    + figure_file = f"figures/{exp_id}.png". Khoá trùng tên cột ở đầu file; giá trị theo đúng danh sách
    chọn của mẫu (CE/MSE, SGD+momentum, ..., diverged Y/N, hidden dạng "256-128", clip "none").
    Chỉ truyền eval_scores (dict có accuracy, macro_f1 — đọc từ eval_result.json) cho baseline và cấu hình cuối."""
    cfg, s = result["cfg"], result["summary"]
    group = GROUP_ALIAS.get(cfg.get("group", "other"), cfg.get("group", "other"))
    if group not in GROUPS:
        group = "other"
    extra = []
    if cfg.get("group") in GROUP_ALIAS:
        extra.append(f"group gốc '{cfg['group']}'")
    if cfg.get("momentum") is not None and cfg["optimizer"] == "sgd_momentum":
        extra.append(f"momentum={cfg['momentum']}")
    if cfg.get("scheduler"):
        extra.append(f"scheduler={cfg['scheduler']}")
    if s.get("diverged"):
        extra.append(f"phân kỳ: {s.get('diverged_at')}")
    if s.get("best_epoch"):
        extra.append(f"best epoch theo val_loss; {s.get('epochs_run')} epoch x {s.get('steps_per_epoch')} bước")
    extra.append("train_loss đo ở eval mode trên tập con cố định 50k mẫu train")
    row = dict(
        exp_id=cfg["exp_id"], group=group, description=cfg.get("description", ""),
        loss=LOSS_NAMES.get(cfg["loss"], cfg["loss"]), optimizer=OPT_NAMES.get(cfg["optimizer"], cfg["optimizer"]),
        lr=cfg["lr"], weight_decay=cfg["weight_decay"], batch=cfg["batch"], epochs=cfg["epochs"],
        hidden="-".join(str(h) for h in cfg["hidden"]), dropout=cfg["dropout"],
        clip_norm="none" if cfg.get("clip_norm") is None else cfg["clip_norm"],
        precision=cfg["precision"], init=cfg["init"], seed=cfg["seed"],
        step0_loss=_num(s.get("step0_loss")), best_val_loss=_num(s.get("best_val_loss")),
        best_epoch=s.get("best_epoch"), final_train_loss=_num(s.get("final_train_loss")),
        final_val_loss=_num(s.get("final_val_loss")), val_acc=_num(s.get("val_acc")),
        val_macro_f1=_num(s.get("val_macro_f1")), time_per_epoch_s=_num(s.get("time_per_epoch_s")),
        peak_mem_MB=_num(s.get("peak_mem_MB")), diverged="Y" if s.get("diverged") else "N",
        eval_acc=None, eval_macro_f1=None, figure_file=f"figures/{cfg['exp_id']}.png",
        notes="; ".join([n for n in [notes] + extra if n]),
    )
    if eval_scores is not None:
        row["eval_acc"] = eval_scores["accuracy"]
        row["eval_macro_f1"] = eval_scores["macro_f1"]
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str,
               seed_ids: list[str] | None = None, summary_notes: dict | None = None) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu, từ dòng 2 trở xuống, rồi lưu thành out_path.

    - Mở mẫu KHÔNG dùng data_only (giữ công thức); đọc tiêu đề dòng 1 để biết cột của từng khoá.
    - Xoá giá trị mẫu ở dòng 2 rồi ghi từng row; BỎ QUA các cột công thức (FORMULA_COLS).
    - seed_ids: exp_id các lần chạy baseline khác seed -> cột A của sheet "Seeds" (dòng 2..6).
    - summary_notes: {group: nhận xét} -> cột "nhận xét ngắn (bạn viết)" của sheet "Summary".
    Sau khi lưu, mở bằng Excel/LibreOffice để công thức tính lại.
    """
    import openpyxl

    assert len(rows) <= N_TEMPLATE_ROWS, f"mẫu chỉ có công thức cho {N_TEMPLATE_ROWS} dòng"
    ids = [r["exp_id"] for r in rows]
    assert len(set(ids)) == len(ids), "exp_id phải duy nhất"

    wb = openpyxl.load_workbook(template_path)
    ws = wb["Experiments"]
    header = {c.value: c.column for c in ws[1] if c.value}
    data_cols = [k for k in header if k not in FORMULA_COLS]
    for k in data_cols:                                  # xoá giá trị mẫu (dòng baseline gợi ý)
        for r in range(2, N_TEMPLATE_ROWS + 2):
            ws.cell(row=r, column=header[k]).value = None
    for i, row in enumerate(rows, start=2):
        unknown = set(row) - set(header)
        assert not unknown, f"khoá không có trong tiêu đề mẫu: {unknown}"
        for k, v in row.items():
            if k in FORMULA_COLS:
                continue
            ws.cell(row=i, column=header[k]).value = _num(v)

    if seed_ids is not None:
        wss = wb["Seeds"]
        assert len(seed_ids) <= 5, "sheet Seeds có chỗ cho tối đa 5 seed (dòng 2..6)"
        for r in range(2, 7):
            wss.cell(row=r, column=1).value = seed_ids[r - 2] if r - 2 < len(seed_ids) else None

    if summary_notes:
        wsm = wb["Summary"]
        hdr = {c.value: c.column for c in wsm[1] if c.value}
        note_col = next(col for name, col in hdr.items() if str(name).startswith("nhận xét"))
        for r in range(2, wsm.max_row + 1):
            g = wsm.cell(row=r, column=1).value
            if g in summary_notes:
                wsm.cell(row=r, column=note_col).value = summary_notes[g]

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
